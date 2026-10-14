"""
分析活頁簿 → 網頁版分析快照（JSON）。

網頁不重寫各分析頁的版面：直接把產好的分析工作表逐格轉成「顯示文字＋樣式」，
數字格式在這裡先轉好（#,##0、0.0%、+/− 帶綠紅、pp），網頁只負責畫表格。
樣式去重成 styles 表；已知底色（標題、分區、斑馬紋、合計、組別色）轉成代號，讓網頁套亮色／暗色。

格式：
{
  "v": 1,
  "styles": [{"b":1, "f":"head", "c":"up", "a":"l", "i":1, "L":1, "T":1}, ...],
  "sheets": [{"name": "...", "widths": [40, 13, ...], "freeze": [列, 欄],
              "rows": [[["文字", 樣式, 跨欄, 跨列] 或 null(被合併蓋住) ...], ...]}]
}
"""
import base64
import gzip
import json
import re
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

FILL_ROLE = {"1F3864": "head", "D6DCE4": "sec", "F2F2F2": "zebra", "DDEBF7": "g1", "E2EFDA": "g2", "FFF2CC": "g3"}
FONT_ROLE = {"FFFFFF": "white", "1F3864": "cur", "2E7D32": "up", "C00000": "down"}


def _hex(color):
    """openpyxl 顏色 → 6 碼大寫 hex；主題色等無法取得時回 None。"""
    try:
        rgb = color.rgb if color is not None else None
    except Exception:
        return None
    if not isinstance(rgb, str) or len(rgb) < 6:
        return None
    return rgb[-6:].upper()


def format_value(value, fmt):
    """依 Excel 數字格式轉成顯示文字；回傳 (文字, 色彩代號 up/down/None)。只處理分析頁用到的格式。"""
    if value is None:
        return "", None
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), None
    if not isinstance(value, (int, float)):
        return str(value), None
    fmt = fmt or "General"
    if fmt in ("General", "@"):
        if isinstance(value, float) and not value.is_integer():
            return f"{value:,.2f}".rstrip("0").rstrip("."), None
        return f"{int(value)}", None
    secs = fmt.split(";")
    if len(secs) > 1:
        if value > 0:
            sec, num = secs[0], value
        elif value < 0:
            sec, num = secs[1], abs(value)
        else:
            sec, num = (secs[2] if len(secs) > 2 else secs[0]), value
    else:
        sec, num = secs[0], value
    color = "up" if "[Color10]" in sec else ("down" if "[Red]" in sec else None)
    sec = re.sub(r"\[[^\]]*\]", "", sec)
    suffix = "".join(re.findall(r'"([^"]*)"', sec))
    sec = re.sub(r'"[^"]*"', "", sec)
    pct = "%" in sec
    m = re.search(r"0\.(0+)", sec)
    decimals = len(m.group(1)) if m else 0
    prefix = sec[0] if sec[:1] in ("+", "-") else ""
    if pct:
        num = num * 100
    q = Decimal(repr(float(num))).quantize(Decimal(1).scaleb(-decimals), ROUND_HALF_UP)   # Excel 四捨五入
    text = f"{q:,.{decimals}f}" if "," in sec else f"{q:.{decimals}f}"
    if text in ("-0", "-0.0", "-0.00"):   # 四捨五入成 0 的負數不顯示負號
        text = text[1:]
    return prefix + text + ("%" if pct else "") + suffix, color


def _style_key(cell, num_color):
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    st = {}
    if font is not None and font.b:
        st["b"] = 1
    fc = _hex(font.color) if font is not None and font.color is not None else None
    if num_color:
        st["c"] = num_color
    elif fc and fc != "000000":
        st["c"] = FONT_ROLE.get(fc, "#" + fc)
    if fill is not None and fill.fill_type == "solid":
        fh = _hex(fill.fgColor)
        if fh and fh != "FFFFFF":
            st["f"] = FILL_ROLE.get(fh, "#" + fh)
    if align is not None:
        if align.horizontal in ("left", "center", "right"):
            st["a"] = align.horizontal[0]
        if align.indent:
            st["i"] = int(align.indent)
        if align.wrap_text:
            st["w"] = 1
    if border is not None:
        if border.left is not None and border.left.style == "medium":
            st["L"] = 1
        if border.top is not None and border.top.style == "medium":
            st["T"] = 1
    return st


def snapshot(wb):
    """分析活頁簿（openpyxl Workbook）→ dict。"""
    from openpyxl.utils import get_column_letter, column_index_from_string
    styles, style_ix = [], {}

    def sid(st):
        key = json.dumps(st, sort_keys=True)
        if key not in style_ix:
            style_ix[key] = len(styles)
            styles.append(st)
        return style_ix[key]

    sheets = []
    for ws in wb.worksheets:
        max_r, max_c = ws.max_row, ws.max_column
        anchors, covered = {}, set()
        for rg in ws.merged_cells.ranges:
            anchors[(rg.min_row, rg.min_col)] = (rg.max_col - rg.min_col + 1, rg.max_row - rg.min_row + 1)
            for r in range(rg.min_row, rg.max_row + 1):
                for c in range(rg.min_col, rg.max_col + 1):
                    if (r, c) != (rg.min_row, rg.min_col):
                        covered.add((r, c))
        rows = []
        for r in range(1, max_r + 1):
            line = []
            for c in range(1, max_c + 1):
                if (r, c) in covered:
                    line.append(None)
                    continue
                cell = ws.cell(row=r, column=c)
                text, num_color = format_value(cell.value, cell.number_format)
                item = [text, sid(_style_key(cell, num_color))]
                span = anchors.get((r, c))
                if span and span != (1, 1):
                    item += [span[0], span[1]]
                line.append(item)
            rows.append(line)
        widths = []
        for c in range(1, max_c + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            widths.append(round(dim.width, 1) if dim is not None and dim.width else 9)
        freeze = [0, 0]
        if ws.freeze_panes:
            m = re.match(r"([A-Z]+)(\d+)", str(ws.freeze_panes))
            if m:
                freeze = [int(m.group(2)) - 1, column_index_from_string(m.group(1)) - 1]
        sheets.append({"name": ws.title, "widths": widths, "freeze": freeze, "rows": rows})
    return {"v": 1, "styles": styles, "sheets": sheets}


def snapshot_payload(analysis_xlsx: bytes):
    """分析 Excel（bytes）→ 上傳用的 (json_gz_b64, xlsx_b64, 原始 JSON 大小)。"""
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(analysis_xlsx))
    raw = json.dumps(snapshot(wb), ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return (base64.b64encode(gzip.compress(raw)).decode(), base64.b64encode(analysis_xlsx).decode(), len(raw))
